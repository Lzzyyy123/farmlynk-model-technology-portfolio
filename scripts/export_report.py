"""Build the public Markdown, DOCX and PDF portfolio deliverables."""
from __future__ import annotations

import re
import subprocess
import textwrap
import tempfile
import zipfile
from xml.sax.saxutils import escape
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
DOCS = [BASE / "docs" / f"{n:02d}-{name}.md" for n, name in [
    (1, "executive-review"), (2, "repository-landscape"), (3, "asset-portfolio"),
    (4, "product-integration"), (5, "remote-sensing"), (6, "land-assessment"),
    (7, "agent-and-knowledge"), (8, "recognition"), (9, "research-results"),
    (10, "intellectual-property"), (11, "topics-and-roadmap"),
    (12, "data-and-evaluation"), (13, "figure-atlas"), (14, "evidence-register"),
    (15, "maintenance"),
]]
OUT = BASE / "deliverables"
OUT.mkdir(exist_ok=True)
TITLE = "技术资产与研发路线 V2.0"


def markdown() -> str:
    chunks = [f"# {TITLE}\n\n公开整理版 · 2026-10-16\n"]
    for path in DOCS:
        text = path.read_text(encoding="utf-8")
        text = re.sub(r"^# ", "## ", text, count=1, flags=re.MULTILINE)
        chunks.append(text.strip())
    return "\n\n---\n\n".join(chunks) + "\n"


def docx_report(md: str) -> None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ModuleNotFoundError:
        _minimal_docx(md)
        return

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    doc.add_heading(TITLE, 0)
    doc.add_paragraph("公开整理版 · 2026-10-16")
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("!["):
            match = re.search(r"\(([^)]+)\)", line)
            if match:
                image = BASE / match.group(1).replace("../", "")
                if image.exists():
                    doc.add_picture(str(image), width=Inches(6.4))
        elif line.startswith("|"):
            continue
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif not line.startswith("```"):
            doc.add_paragraph(line)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "STHeiti"
            run.font.size = Pt(10)
    doc.save(OUT / f"{TITLE}.docx")


def _minimal_docx(md: str) -> None:
    """Write a readable DOCX using only the standard library when needed."""
    paragraphs = []
    for raw in md.splitlines():
        line = raw.strip()
        if not line or line.startswith("|") or line.startswith("---") or line.startswith("!"):
            continue
        line = re.sub(r"[`*_]", "", line)
        if line.startswith("#"):
            line = line.lstrip("# ")
        paragraphs.append(line)
    body = []
    for line in paragraphs:
        body.append("<w:p><w:r><w:t xml:space=\"preserve\">%s</w:t></w:r></w:p>" % escape(line))
    document = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
        "<w:body>%s<w:sectPr><w:pgSz w:w=\"11906\" w:h=\"16838\"/>"
        "<w:pgMar w:top=\"900\" w:right=\"900\" w:bottom=\"900\" w:left=\"900\"/>"
        "</w:sectPr></w:body></w:document>" % "".join(body)
    )
    content_types = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
        "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
        "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
        "<Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
        "</Types>"
    )
    rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/>"
        "</Relationships>"
    )
    with zipfile.ZipFile(OUT / f"{TITLE}.docx", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)


def pdf_report(md: str) -> None:
    try:
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.font_manager import FontProperties
    except ModuleNotFoundError:
        _cups_pdf(md)
        return

    font_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    prop = FontProperties(fname=font_path) if Path(font_path).exists() else None
    lines = []
    for raw in md.splitlines():
        line = re.sub(r"[`*_]", "", raw).strip()
        if not line or line.startswith("|") or line.startswith("---"):
            continue
        if line.startswith("!"):
            continue
        lines.extend(textwrap.wrap(line, width=54) or [""])
    with PdfPages(OUT / f"{TITLE}.pdf") as pdf:
        page = []
        for line in lines:
            page.append(line)
            if len(page) == 38:
                _pdf_page(pdf, page, prop)
                page = []
        if page:
            _pdf_page(pdf, page, prop)


def _cups_pdf(md: str) -> None:
    """Use the macOS text-to-PDF filter when plotting libraries are absent."""
    lines = []
    for raw in md.splitlines():
        line = re.sub(r"[`*_]", "", raw).strip()
        if not line or line.startswith("|") or line.startswith("---") or line.startswith("!"):
            continue
        lines.extend(textwrap.wrap(line, width=72) or [""])
    with tempfile.NamedTemporaryFile("w", suffix=".txt", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
        handle.flush()
        with (OUT / f"{TITLE}.pdf").open("wb") as output:
            result = subprocess.run(["cupsfilter", "-m", "application/pdf", handle.name], stdout=output, stderr=subprocess.PIPE)
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode("utf-8", "replace"))


def _pdf_page(pdf, lines, prop) -> None:
    fig = plt.figure(figsize=(8.27, 11.69), facecolor="white")
    fig.text(0.07, 0.96, TITLE, fontsize=16, fontproperties=prop, color="#17313B")
    y = 0.92
    for line in lines:
        size = 11 if line.startswith("#") else 9
        fig.text(0.07, y, line, fontsize=size, fontproperties=prop, color="#24343A", va="top")
        y -= 0.022 if size == 9 else 0.028
    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    md = markdown()
    (OUT / f"{TITLE}.md").write_text(md, encoding="utf-8")
    docx_report(md)
    pdf_report(md)
    print({"deliverables": 3, "directory": str(OUT)})


if __name__ == "__main__":
    main()
